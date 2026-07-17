from __future__ import annotations

import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.schemas.export import (
    ExcelExportRequest,
)
from app.services.excel_columns import (
    EXPORT_COLUMN_CONFIG,
)
from app.services.listing_export_query import (
    build_listing_export_query,
    stream_listing_rows,
)


VN_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")


HEADER_FONT = Font(
    bold=True,
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

HEADER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
)

DATA_ALIGNMENT = Alignment(
    vertical="top",
)


def sanitize_filename(
    filename: str,
) -> str:
    filename = re.sub(
        r'[<>:"/\\|?*\x00-\x1F]',
        "_",
        filename,
    )

    filename = filename.strip(" .")

    if not filename:
        filename = "marketplace-listings"

    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"

    return filename[:180]


def normalize_excel_value(
    value: Any,
) -> Any:
    if value is None:
        return None

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(VN_TIMEZONE)

            value = value.replace(tzinfo=None)

        return value

    if isinstance(
        value,
        (str, int, float, bool),
    ):
        return value

    return str(value)


def configure_sheet(
    worksheet,
    columns: list[str],
):
    worksheet.freeze_panes = "A2"

    for index, column in enumerate(
        columns,
        start=1,
    ):
        config = EXPORT_COLUMN_CONFIG[column]

        column_letter = get_column_letter(index)

        worksheet.column_dimensions[column_letter].width = config.get(
            "width",
            18,
        )


def append_header(
    worksheet,
    columns: list[str],
):
    header_cells = []

    for column in columns:
        config = EXPORT_COLUMN_CONFIG[column]

        cell = WriteOnlyCell(
            worksheet,
            value=config["label"],
        )

        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT

        header_cells.append(cell)

    worksheet.append(header_cells)


def append_listing_row(
    worksheet,
    listing,
    columns: list[str],
):
    row_cells = []

    for column in columns:
        config = EXPORT_COLUMN_CONFIG[column]

        raw_value = getattr(
            listing,
            column,
            None,
        )

        value = normalize_excel_value(raw_value)

        cell = WriteOnlyCell(
            worksheet,
            value=value,
        )

        cell.alignment = DATA_ALIGNMENT

        number_format = config.get("format")

        if number_format:
            cell.number_format = number_format

        if (
            column == "listing_url"
            and isinstance(value, str)
            and value.startswith(("http://", "https://"))
        ):
            cell.hyperlink = value
            cell.style = "Hyperlink"

        row_cells.append(cell)

    worksheet.append(row_cells)


def create_export_info_sheet(
    workbook: Workbook,
    request: ExcelExportRequest,
    total_rows: int,
):
    worksheet = workbook.create_sheet("Export Info")

    filter_data = request.filters

    info_rows = [
        (
            "Export time",
            datetime.now(VN_TIMEZONE).strftime("%d/%m/%Y %H:%M:%S"),
        ),
        (
            "Scope",
            request.scope,
        ),
        (
            "Total rows",
            total_rows,
        ),
        (
            "Search",
            filter_data.q or "",
        ),
        (
            "Marketplaces",
            ", ".join(filter_data.marketplaces),
        ),
        (
            "Statuses",
            ", ".join(filter_data.statuses),
        ),
        (
            "Sync action",
            filter_data.action or "",
        ),
        (
            "Product ID",
            filter_data.product_id or "",
        ),
        (
            "Brand",
            filter_data.brand or "",
        ),
        (
            "Category",
            filter_data.category_name or "",
        ),
        (
            "Condition",
            filter_data.condition or "",
        ),
        (
            "Minimum price",
            filter_data.price_min,
        ),
        (
            "Maximum price",
            filter_data.price_max,
        ),
        (
            "Sort",
            (f"{request.sort.field} {request.sort.order}"),
        ),
        (
            "Columns",
            ", ".join(request.columns),
        ),
    ]

    for label, value in info_rows:
        label_cell = WriteOnlyCell(
            worksheet,
            value=label,
        )

        label_cell.font = Font(bold=True)

        value_cell = WriteOnlyCell(
            worksheet,
            value=normalize_excel_value(value),
        )

        worksheet.append(
            [
                label_cell,
                value_cell,
            ]
        )

    worksheet.column_dimensions["A"].width = 24

    worksheet.column_dimensions["B"].width = 80


def export_listings_to_excel(
    db: Session,
    request: ExcelExportRequest,
) -> tuple[Path, str, int]:
    query = build_listing_export_query(
        filters=request.filters,
        sort=request.sort,
    )

    if request.scope == "CURRENT_PAGE":
        offset = (request.page - 1) * request.page_size

        query = query.offset(offset).limit(request.page_size)

    default_filename = (
        f"marketplace-listings-{datetime.now(VN_TIMEZONE):%Y%m%d-%H%M%S}.xlsx"
    )

    filename = sanitize_filename(request.filename or default_filename)

    temporary_file = NamedTemporaryFile(
        prefix="marketplace-export-",
        suffix=".xlsx",
        delete=False,
    )

    temporary_path = Path(temporary_file.name)

    temporary_file.close()

    workbook = Workbook(write_only=True)

    listing_sheet = workbook.create_sheet("Listings")

    configure_sheet(
        listing_sheet,
        request.columns,
    )

    append_header(
        listing_sheet,
        request.columns,
    )

    exported_rows = 0

    try:
        for listing in stream_listing_rows(
            db=db,
            query=query,
            batch_size=2000,
        ):
            append_listing_row(
                worksheet=listing_sheet,
                listing=listing,
                columns=request.columns,
            )

            exported_rows += 1

        if exported_rows == 0:
            raise ValueError("Không có dữ liệu phù hợp để export")

        create_export_info_sheet(
            workbook=workbook,
            request=request,
            total_rows=exported_rows,
        )

        workbook.save(temporary_path)

        return (
            temporary_path,
            filename,
            exported_rows,
        )

    except Exception:
        temporary_path.unlink(missing_ok=True)

        raise
